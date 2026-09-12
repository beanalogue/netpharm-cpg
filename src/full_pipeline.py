"""
전체 자동 분석 파이프라인 오케스트레이터
AI 설정 탭 원클릭 실행용 — 각 단계 진행 상황을 Queue로 전달
"""

import sys
import time
import logging
import queue
from pathlib import Path

BASE_DIR = Path(__file__).parent.parent
sys.path.insert(0, str(BASE_DIR / "src"))
log = logging.getLogger(__name__)


def _tg_send(text: str):
    """텔레그램 알림 (실패해도 무시)"""
    try:
        from telegram_notify import send
        send(text)
    except Exception:
        pass


def _tg_send_chunks(text: str, header: str = ""):
    """긴 텍스트를 4000자씩 나눠서 전송"""
    MAX = 4000
    if header:
        _tg_send(header)
    for i in range(0, len(text), MAX):
        chunk = text[i:i + MAX]
        _tg_send(chunk)


def run_full_analysis(
    cfg: dict,
    q: "queue.Queue",
    api_key: str = "",
    do_docking: bool = False,
    do_ai: bool = True,
):
    """
    전체 파이프라인 순차 실행.

    큐 메시지 형식:
      {"type": "progress", "step": N, "total": T, "msg": str, "done": bool, "warn": bool}
      {"type": "log",      "msg": str}
      {"type": "done",     "result": dict}
      {"type": "error",    "msg": str}
    """
    TOTAL = sum([6, do_docking, do_ai])   # 기본 6단계 + 선택 단계
    result = {}

    def emit(step, msg, done=False, warn=False):
        q.put({"type": "progress", "step": step, "total": TOTAL,
               "msg": msg, "done": done, "warn": warn})

    def logmsg(msg):
        q.put({"type": "log", "msg": msg})

    # ─────────────────────────────────────────────────────────────
    # Step 1  화합물·타깃 수집
    # ─────────────────────────────────────────────────────────────
    emit(1, "화합물·타깃 수집 중...")
    try:
        import sys; sys.path.insert(0, str(BASE_DIR / "src"))
        from pipeline import run_pipeline, DEFAULT_CONFIG
        t0 = time.time()
        res = run_pipeline({**DEFAULT_CONFIG, **cfg})
        res["config"]  = cfg
        res["elapsed"] = time.time() - t0
        result["pipeline"] = res
        n_genes = len(res.get("genes", []))

        # 타깃 0개 약재 감지 → 경고 로그
        zero_target_herbs = []
        for herb_name, hd in res.get("herb_results", {}).items():
            total_targets = sum(len(c.get("targets", [])) for c in hd.get("compounds", []))
            if total_targets == 0:
                zero_target_herbs.append(herb_name)
        if zero_target_herbs:
            warn_msg = f"타깃 0개 약재: {', '.join(zero_target_herbs)} — DB 검증 필요"
            logmsg(f"[경고] {warn_msg}")
            emit(1, f"타깃 수집 완료 — 유전자 {n_genes}개 ({res['elapsed']:.0f}초) ⚠ {warn_msg}",
                 done=True, warn=True)
            result["zero_target_herbs"] = zero_target_herbs
        else:
            emit(1, f"타깃 수집 완료 — 유전자 {n_genes}개 ({res['elapsed']:.0f}초)", done=True)
    except Exception as e:
        _tg_send(f"❌ <b>[{cfg.get('name','')}] 수집 오류</b>\n{e}")
        q.put({"type": "error", "msg": f"[수집 오류] {e}"})
        return

    # ─────────────────────────────────────────────────────────────
    # Step 2  ADMET 스크리닝
    # ─────────────────────────────────────────────────────────────
    emit(2, "ADMET 약물성 스크리닝 중...")
    admet_res = None
    n_pass    = 0
    screener  = None
    try:
        from admet import ADMETScreener, ADMETCriteria
        admet_cfg = cfg.get("admet", {})
        criteria  = ADMETCriteria(**{k: v for k, v in admet_cfg.items()
                                      if k in ADMETCriteria.__dataclass_fields__})
        screener  = ADMETScreener(criteria)
        admet_res = screener.screen_all(res["herb_results"])
        result["admet"] = admet_res
        n_pass = len(admet_res.get("passed", []))
        n_fail = len(admet_res.get("failed", []))
        emit(2, f"ADMET 완료 — {n_pass}개 통과 / {n_fail}개 탈락", done=True,
             warn=(n_pass == 0))
    except Exception as e:
        logmsg(f"[ADMET 경고] {e}")
        _tg_send(f"⚠️ <b>[{cfg.get('name','')}] ADMET 실패</b>\n{e}")
        emit(2, "ADMET 실패 (건너뜀)", done=True, warn=True)

    # ─────────────────────────────────────────────────────────────
    # Step 3  ADMET 필터 네트워크 재구성
    # ─────────────────────────────────────────────────────────────
    emit(3, "ADMET 통과 성분 기반 네트워크 재구성 중...")
    active_res = res
    try:
        if admet_res and n_pass > 0:
            filt_hr = screener.filter_herb_results(res["herb_results"])
            genes   = sorted({
                t["gene_symbol"].upper()
                for hd in filt_hr.values()
                for c  in hd.get("compounds", [])
                for t  in c.get("targets", [])
                if t.get("gene_symbol")
            })
            from collect import NetPharmCollector
            from network   import build_full_network
            from enrichment import run_enrichment
            nc   = NetPharmCollector()
            ppi  = nc.collect_ppi(genes, min_score=cfg.get("ppi_min_score", 400))
            nc.close()
            name_a  = cfg.get("name", "analysis") + "_admet"
            net_f   = build_full_network(filt_hr, ppi, prefix=name_a,
                                          min_ppi_score=cfg.get("ppi_min_score", 400))
            enr_f   = run_enrichment(genes, prefix=name_a,
                                      adj_p_cutoff=cfg.get("adj_p_cutoff", 0.05),
                                      top_n=cfg.get("top_n_enrichment", 20),
                                      make_plots=True)
            active_res = {**res, "herb_results": filt_hr, "genes": genes,
                           "network": net_f, "enrichment": enr_f,
                           "config": {**cfg, "name": name_a}}
            result["admet_filtered"] = active_res
            emit(3, f"재구성 완료 — {len(genes)}개 유전자 / PPI {len(ppi)}개 엣지", done=True)
        else:
            emit(3, "통과 성분 없음 — 전체 데이터로 계속", done=True, warn=True)
    except Exception as e:
        logmsg(f"[Step 3 경고] {e}")
        _tg_send(f"⚠️ <b>[{cfg.get('name','')}] 네트워크 재구성 실패</b>\n{e}")
        emit(3, "재구성 실패 (전체 데이터 사용)", done=True, warn=True)

    # ─────────────────────────────────────────────────────────────
    # Step 4  질환 교차분석
    # ─────────────────────────────────────────────────────────────
    disease_q = cfg.get("disease_query", "")
    if disease_q:
        emit(4, f"질환 타깃 조회 중: '{disease_q}'...")
        try:
            from disease import DiseaseTargetCollector, IntersectionAnalyzer, IntersectionPlotter
            from collect  import CacheDB
            cache    = CacheDB()
            col_dis  = DiseaseTargetCollector(cache)
            _, d_tgt = col_dis.get_targets_by_name(
                disease_q, min_score=cfg.get("disease_min_score", 0.1))
            cache.close()
            herb_genes = set(active_res.get("genes", []))
            d_genes    = {t["gene_symbol"] for t in d_tgt}
            inters     = IntersectionAnalyzer().intersect(
                herb_genes, d_genes, herb_label="Herbs", disease_label=disease_q)
            IntersectionPlotter(BASE_DIR / "results").venn_diagram(
                inters, prefix=cfg.get("name", "analysis"))
            result["intersection"]    = inters
            result["disease_targets"] = d_tgt
            emit(4, f"교차분석 완료 — 공유 타깃 {inters.get('shared_count', 0)}개", done=True)
        except Exception as e:
            logmsg(f"[Step 4 경고] {e}")
            _tg_send(f"⚠️ <b>[{cfg.get('name','')}] 질환교차 실패</b>\n{e}")
            emit(4, "질환교차 실패 (건너뜀)", done=True, warn=True)
    else:
        emit(4, "질환 쿼리 없음 — 건너뜀", done=True)

    # ─────────────────────────────────────────────────────────────
    # Step 5  허브 유전자 토폴로지
    # ─────────────────────────────────────────────────────────────
    emit(5, "허브 유전자 토폴로지 분석 중 (Degree / Betweenness / MCC)...")
    topo_df = None
    try:
        net = active_res.get("network", {})
        # build_full_network이 이미 topology 계산해서 반환함 — 재사용
        topo_df = net.get("topology") if net.get("topology") is not None else None
        if topo_df is None or topo_df.empty:
            # 없으면 직접 계산 (ADMET 재구성 경로 등)
            from topology import TopologyAnalyzerExtended
            net_obj = net.get("integrated") or net.get("ppi")
            if net_obj is not None:
                raw_G = getattr(net_obj, "G", net_obj)  # .G 있으면 사용, 없으면 그대로
                ta      = TopologyAnalyzerExtended(raw_G)
                topo_df = ta.compute(node_type_filter="target")
        # compound 노드 제거 — target 노드만 허브 분석에 사용
        if topo_df is not None and not topo_df.empty:
            if "node_type" in topo_df.columns:
                topo_df = topo_df[topo_df["node_type"] == "target"].reset_index(drop=True)
            result["topology"] = topo_df
            top5 = ", ".join(topo_df.head(5)["label"].tolist())
            emit(5, f"토폴로지 완료 — Top 허브: {top5}", done=True)
        else:
            emit(5, "네트워크 없음 — 건너뜀", done=True, warn=True)
    except Exception as e:
        logmsg(f"[Step 5 경고] {e}")
        emit(5, "토폴로지 실패 (건너뜀)", done=True, warn=True)

    # ─────────────────────────────────────────────────────────────
    # Step 6 (선택)  분자 도킹
    # ─────────────────────────────────────────────────────────────
    step_offset = 0
    if do_docking:
        step_offset = 1
        emit(6, "분자 도킹 실행 중 (허브 Top3 × 상위 성분 Top5) — 수 분 소요...")
        try:
            from docking import BatchDocking, check_tools
            tools = check_tools()
            if tools["vina"] and tools["obabel"] and topo_df is not None:
                # 리간드: ADMET 통과 성분 우선, 없으면 전체 성분에서 SMILES 있는 것
                if admet_res and admet_res.get("passed"):
                    passed = admet_res["passed"]
                    ligs = [{"chembl_id": r.compound_id,
                             "pref_name": r.compound_name,
                             "smiles":    r.smiles}
                            for r in passed[:5] if getattr(r, "smiles", "")]
                else:
                    # ADMET 없으면 herb_results에서 SMILES 추출
                    ligs = []
                    for hd in res.get("herb_results", {}).values():
                        for c in hd.get("compounds", []):
                            mol = c.get("molecule", {})
                            smi = mol.get("smiles", "")
                            cid = mol.get("chembl_id", "")
                            nm  = mol.get("pref_name", cid)
                            if smi and cid and len(ligs) < 5:
                                ligs.append({"chembl_id": cid, "pref_name": nm, "smiles": smi})

                # 수용체: 토폴로지 Top3 허브 유전자 (compound 노드 제외)
                target_rows = topo_df[topo_df.get("node_type", topo_df.get("type", "")) == "target"] \
                    if "node_type" in topo_df.columns or "type" in topo_df.columns \
                    else topo_df
                recs = [{"label":      str(row.get("label", "")),
                         "uniprot_id": str(row.get("uniprot_id", "") or "")}
                        for _, row in target_rows.head(3).iterrows()
                        if row.get("label")]

                if ligs and recs:
                    batcher = BatchDocking(max_ligands=5, max_receptors=3)
                    dock_df = batcher.run(ligs, recs)
                    result["docking"] = dock_df
                    n_ok = int((dock_df["affinity"].notna()).sum()) if dock_df is not None and not dock_df.empty else 0
                    # 도킹 히트맵 생성
                    if dock_df is not None and not dock_df.empty:
                        try:
                            from figure_composer import make_docking_heatmap
                            _hm = make_docking_heatmap(dock_df, prefix=cfg.get("name","analysis"),
                                                        output_dir=BASE_DIR/"results")
                            result["docking_heatmap"] = str(_hm) if _hm else ""
                        except Exception as _de:
                            logmsg(f"[도킹 히트맵 경고] {_de}")
                    emit(6, f"도킹 완료 — {n_ok}/{len(ligs)*len(recs)}쌍 성공", done=True, warn=(n_ok == 0))
                else:
                    emit(6, f"도킹 입력 부족 (리간드 {len(ligs)}개, 수용체 {len(recs)}개) — 건너뜀", done=True, warn=True)
            else:
                missing = "Vina" if not tools["vina"] else "OpenBabel"
                emit(6, f"{missing} 미설치 또는 네트워크 없음 — 건너뜀", done=True, warn=True)
        except Exception as e:
            logmsg(f"[도킹 경고] {e}")
            emit(6, "도킹 실패 (건너뜀)", done=True, warn=True)

    # ─────────────────────────────────────────────────────────────
    # Step 6/7 (선택)  AI 논문 초안
    # ─────────────────────────────────────────────────────────────
    if do_ai and api_key:
        ai_step = 6 + step_offset
        emit(ai_step, "AI 논문 초안 생성 중 (Methods / Results / Discussion)...")
        try:
            from interpreter import (interpret_network, draft_methods, draft_results,
                                      draft_discussion, draft_front_matter, _make_figure_legends)
            import time as _t
            net_r    = active_res.get("network", {})
            enr_r    = active_res.get("enrichment", {})
            cfg_     = active_res.get("config", cfg)
            herbs    = list((cfg_.get("herbs") or {}).keys())
            drugs    = cfg_.get("drugs", [])
            hubs_df  = topo_df
            hub_genes = (hubs_df[["label", "degree", "betweenness"]].to_dict("records")
                         if hubs_df is not None and not hubs_df.empty else [])
            er       = (enr_r.get("results") or {})
            kegg_df  = er.get("KEGG")
            gobp_df  = er.get("GO_BP")
            kegg     = (kegg_df.head(10)[["term", "gene_count", "adj_p_value"]].to_dict("records")
                        if kegg_df is not None and not kegg_df.empty else [])
            gobp     = (gobp_df.head(8)[["term", "gene_count", "adj_p_value"]].to_dict("records")
                        if gobp_df is not None and not gobp_df.empty else [])
            hct = net_r.get("hct"); ppi_n = net_r.get("ppi")
            stats = dict(
                hct_nodes=hct.G.number_of_nodes() if hct else 0,
                hct_edges=hct.G.number_of_edges() if hct else 0,
                ppi_nodes=ppi_n.G.number_of_nodes() if ppi_n else 0,
                ppi_edges=ppi_n.G.number_of_edges() if ppi_n else 0,
                total_genes=len(active_res.get("genes", [])),
            )
            herb_dict = cfg_.get("herbs", {})

            # 5 RPM 제한 준수: 각 호출 사이 13초
            ai_out = {}
            ai_out["interpretation"] = interpret_network(herbs, drugs, hub_genes, kegg, gobp, api_key)
            _t.sleep(13)
            ai_out["methods"]    = draft_methods(herbs, drugs, herb_dict, api_key=api_key)
            _t.sleep(13)
            ai_out["results"]    = draft_results(herbs, drugs, stats, hub_genes, kegg, api_key)
            _t.sleep(13)
            ai_out["discussion"] = draft_discussion(herbs, drugs, hub_genes, kegg,
                                                     ai_out["interpretation"], api_key)
            _t.sleep(13)
            front   = draft_front_matter(herbs, drugs, hub_genes, kegg, stats,
                                         ai_out["interpretation"], api_key)
            ai_out["abstract"]     = front.get("abstract", "")
            ai_out["introduction"] = front.get("introduction", "")
            ai_out["conclusion"]   = front.get("conclusion", "")
            ai_out["figure_legends"] = _make_figure_legends(herbs, drugs, stats, hub_genes, kegg)
            result["ai_output"] = ai_out

            # Markdown 저장
            name = cfg.get("name", "analysis")
            md_sections = [
                ("abstract",       "## Abstract"),
                ("introduction",   "## Introduction"),
                ("methods",        "## Methods"),
                ("results",        "## Results"),
                ("discussion",     "## Discussion"),
                ("conclusion",     "## Conclusion"),
                ("figure_legends", "## Figure Legends"),
                ("interpretation", "## 작용기전 해석 (Korean)"),
            ]
            md_lines = [f"# {name} — AI-Generated Manuscript Draft\n"]
            for sec, lbl in md_sections:
                if ai_out.get(sec):
                    md_lines.append(f"{lbl}\n\n{ai_out[sec]}\n")
            md_path = BASE_DIR / "results" / f"{name}_ai_draft.md"
            md_path.write_text("\n".join(md_lines), encoding="utf-8")
            result["ai_draft_path"] = str(md_path)

            # Word(.docx) 저장
            try:
                from docx import Document
                from docx.shared import Pt
                doc = Document()
                doc.add_heading(f"{name} — Network Pharmacology Manuscript", 0)
                for sec, lbl in md_sections:
                    txt = ai_out.get(sec, "")
                    if not txt:
                        continue
                    doc.add_heading(lbl.lstrip("# "), level=1)
                    for para in txt.split("\n\n"):
                        para = para.strip()
                        if para:
                            doc.add_paragraph(para)
                docx_path = BASE_DIR / "results" / f"{name}_manuscript.docx"
                doc.save(str(docx_path))
                result["manuscript_docx"] = str(docx_path)
            except Exception as _de:
                logmsg(f"[Word 저장 경고] {_de}")

            emit(ai_step, "AI 초안 완료 (Abstract / Intro / Methods / Results / Discussion / Conclusion)", done=True)
        except Exception as e:
            logmsg(f"[AI 경고] {e}")
            _tg_send(f"⚠️ <b>[{cfg.get('name','')}] AI 초안 실패</b>\n{e}")
            emit(ai_step, "AI 초안 생성 실패 (건너뜀)", done=True, warn=True)

    # ─────────────────────────────────────────────────────────────
    # 마지막  Excel + 통합 Figure 내보내기
    # ─────────────────────────────────────────────────────────────
    last_step = TOTAL
    emit(last_step, "Excel + 통합 Figure 생성 중...")
    try:
        from exporter       import SupplementaryExporter, MetadataLogger
        from figure_composer import make_publication_figure

        name    = cfg.get("name", "analysis")
        meta    = MetadataLogger()
        supp    = SupplementaryExporter()
        xl_path = supp.build_excel(
            prefix          = name,
            admet_df        = (admet_res or {}).get("summary"),
            topology_df     = topo_df,
            enrich_results  = (active_res.get("enrichment") or {}).get("results"),
            intersection    = result.get("intersection"),
            disease_targets = result.get("disease_targets"),
            metadata        = meta,
        )
        result["excel_path"] = str(xl_path) if xl_path else ""

        re_dir = BASE_DIR / "results"
        en_dir = BASE_DIR / "enrichment"

        def _find(folder, pats):
            for p in pats:
                hits = sorted(folder.glob(p))
                if hits: return str(hits[0])
            return None

        fig_out = make_publication_figure(
            venn_path    = _find(re_dir, [f"{name}_venn.png"]),
            hub_path     = _find(re_dir, [f"{name}*hub_ranking.png"]),
            dotplot_path = _find(en_dir, [f"{name}*combined_dotplot.png"]),
            docking_path = result.get("docking_heatmap") or _find(re_dir, [f"{name}_docking_heatmap.png"]),
            prefix       = name,
            dpi          = 300,
            output_dir   = re_dir,
        )
        result["figure_png"] = str(fig_out.get("png", ""))
        result["figure_svg"] = str(fig_out.get("svg", ""))

        emit(last_step, "파일 생성 완료", done=True)
    except Exception as e:
        logmsg(f"[내보내기 경고] {e}")
        emit(last_step, "일부 파일 생성 실패", done=True, warn=True)

    # ── 완료 알림: AI 해석 텍스트 + 분석 요약 전송 ──────────────
    name = cfg.get("name", "analysis")
    ai_out = result.get("ai_output", {})
    pipe   = result.get("pipeline", {})

    # 분석 요약 (항상)
    genes  = pipe.get("genes", [])
    net    = pipe.get("network", {})
    hubs   = result.get("topology")
    hub_txt = ""
    if hubs is not None and not hubs.empty:
        hub_txt = ", ".join(hubs.head(5)["label"].tolist())
    enrich  = (pipe.get("enrichment") or {}).get("results", {})
    kegg_df = enrich.get("KEGG")
    kegg_txt = ""
    if kegg_df is not None and not kegg_df.empty:
        kegg_txt = "\n".join(
            f"  • {r['term'][:50]} (p={r['adj_p_value']:.1e})"
            for _, r in kegg_df.head(5).iterrows()
        )

    summary = (
        f"✅ <b>[{name}] 분석 완료!</b>\n"
        f"• 유전자: {len(genes)}개\n"
        f"• Hub 타겟: {hub_txt or '—'}\n"
        f"• Top KEGG:\n{kegg_txt or '  —'}"
    )
    _tg_send(summary)

    # AI 작용기전 해석 (한국어) — 있을 때만
    interp = ai_out.get("interpretation", "")
    if interp:
        _tg_send_chunks(interp, header=f"🔬 <b>[{name}] AI 작용기전 해석</b>")

    # Abstract (영어)
    abstract = ai_out.get("abstract", "")
    if abstract:
        _tg_send_chunks(abstract, header=f"📄 <b>[{name}] Abstract</b>")

    # Publication Guide
    import glob as _gl
    re_dir  = BASE_DIR / "results"
    en_dir  = BASE_DIR / "enrichment"
    net_dir = BASE_DIR / "networks"
    def _ex(folder, pat): return bool(_gl.glob(str(folder / pat)))

    main_lines = []
    supp_lines = []
    if _ex(re_dir, f"{name}_figure_multipanel*"):
        main_lines.append(f"Fig 1: Multipanel figure")
    else:
        if _ex(re_dir, f"{name}_network*"): main_lines.append("Fig 1: PPI/Integrated network")
        if _ex(re_dir, f"{name}_hub_ranking*"): main_lines.append("Fig 2: Hub gene ranking")
        if _ex(re_dir, f"{name}_venn*"): main_lines.append("Fig 3: Venn diagram")
        if _ex(en_dir, f"{name}_KEGG*dotplot*"): main_lines.append("Fig 4: KEGG dotplot")
        if _ex(re_dir, f"{name}_docking_heatmap*"): main_lines.append("Fig 5: Docking heatmap")
    if _ex(re_dir, f"{name}_admet_network*"): supp_lines.append("Supp Fig S1: ADMET network")
    if _ex(re_dir, f"{name}_admet_hub_ranking*"): supp_lines.append("Supp Fig S2: ADMET hub ranking")
    for db in ["GO_BP","GO_MF","GO_CC","Reactome"]:
        if _ex(en_dir, f"{name}_{db}*dotplot*"): supp_lines.append(f"Supp Fig: {db} dotplot")
    if _ex(re_dir, f"{name}_supplementary*"): supp_lines.append("Supp Table S1: 전체 데이터 Excel")

    guide = (
        f"📋 <b>[{name}] Publication Guide</b>\n\n"
        f"<b>Main Figures (본문)</b>\n"
        + ("\n".join(f"  • {l}" for l in main_lines) or "  —") + "\n\n"
        f"<b>Main Tables (본문)</b>\n"
        f"  • Table 1: 약재-성분-타깃 요약\n"
        f"  • Table 2: Hub 타겟 Top 10~20\n"
        f"  • Table 3: Top KEGG 경로\n\n"
        f"<b>Supplementary (부록)</b>\n"
        + ("\n".join(f"  • {l}" for l in supp_lines) or "  —") + "\n\n"
        f"<b>참고용 (미포함)</b>\n"
        f"  • *.graphml (Cytoscape용)\n"
        f"  • _ai_draft.md / _manuscript.docx (AI 초안)"
    )
    _tg_send(guide)

    q.put({"type": "done", "result": result})
